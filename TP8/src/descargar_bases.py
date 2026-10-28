"""Ingesta y preparación de bases del TP8. No produce gráficos ni respuestas."""
import argparse
import csv
import hashlib
import io
import json
import math
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from time import sleep
from urllib.request import Request, urlopen

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw'
OUT = ROOT / 'data/processed'
COUNTRIES = {'IND': 'India', 'USA': 'Estados Unidos'}
# Alias, código WDI, unidad, bloque de la consigna.
SERIES = [
    ('natalidad_por_mil', 'SP.DYN.CBRT.IN', 'nacimientos por 1.000 habitantes', 'demografia'),
    ('mortalidad_por_mil', 'SP.DYN.CDRT.IN', 'muertes por 1.000 habitantes', 'demografia'),
    ('poblacion_total', 'SP.POP.TOTL', 'personas', 'demografia'),
    ('crecimiento_poblacional_pct_anual', 'SP.POP.GROW', '% anual; crecimiento total de la población', 'demografia'),
    ('poblacion_0_14', 'SP.POP.0014.TO', 'personas', 'demografia'),
    ('poblacion_15_64', 'SP.POP.1564.TO', 'personas', 'demografia'),
    ('poblacion_65_mas', 'SP.POP.65UP.TO', 'personas', 'demografia'),
    ('dependencia_total_wdi', 'SP.POP.DPND', 'dependientes por 100 personas de 15-64', 'demografia'),
    ('dependencia_jovenes_wdi', 'SP.POP.DPND.YG', 'menores de 15 por 100 personas de 15-64', 'demografia'),
    ('dependencia_mayores_wdi', 'SP.POP.DPND.OL', 'personas de 65+ por 100 personas de 15-64', 'demografia'),
    ('fecundidad', 'SP.DYN.TFRT.IN', 'hijos por mujer', 'demografia'),
    ('esperanza_vida', 'SP.DYN.LE00.IN', 'años', 'demografia'),
    ('migracion_neta', 'SM.POP.NETM', 'personas; inmigrantes menos emigrantes durante el periodo', 'migracion'),
    ('remesas_recibidas_usd', 'BX.TRF.PWKR.CD.DT', 'US$ corrientes', 'divisas'),
    ('remesas_recibidas_pct_pib_wdi', 'BX.TRF.PWKR.DT.GD.ZS', '% del PIB', 'divisas'),
    ('pib_usd', 'NY.GDP.MKTP.CD', 'US$ corrientes', 'divisas'),
    ('fbkf_usd', 'NE.GDI.FTOT.CD', 'US$ corrientes', 'divisas'),
    ('fbkf_pct_pib_wdi', 'NE.GDI.FTOT.ZS', '% del PIB', 'divisas'),
    ('exportaciones_usd', 'NE.EXP.GNFS.CD', 'US$ corrientes', 'divisas'),
    ('exportaciones_pct_pib_wdi', 'NE.EXP.GNFS.ZS', '% del PIB', 'divisas'),
    ('ied_entrada_neta_usd', 'BX.KLT.DINV.CD.WD', 'US$ corrientes; entradas netas', 'divisas'),
    ('ied_entrada_neta_pct_pib_wdi', 'BX.KLT.DINV.WD.GD.ZS', '% del PIB', 'divisas'),
    ('deuda_externa_flujo_neto_usd', 'DT.NFL.DECT.CD', 'US$ corrientes; flujo neto', 'divisas'),
    ('deuda_largo_plazo_desembolsos_usd', 'DT.DIS.DLXF.CD', 'US$ corrientes; desembolsos brutos', 'divisas'),
    ('aod_neta_recibida_usd', 'DT.ODA.ODAT.CD', 'US$ corrientes; ayuda neta recibida', 'divisas'),
]
HDR_BASE = 'https://hdr.undp.org/sites/default/files/2025_HDR/'
HDR_FILES = [
    'HDR25_Composite_indices_complete_time_series.csv',
    'HDR25_Composite_indices_metadata.xlsx',
    'HDR25_Statistical_Annex_GII_Table.xlsx',
    'HDR25_Technical_Notes.pdf',
]
IDS_CODES = {'DT.NFL.DECT.CD', 'DT.DIS.DLXF.CD'}
IDS_HISTORICAL_END = 2024  # IDS publicado el 03/12/2025; datos históricos hasta 2024.


def series_file(code):
    return f"{'ids' if code in IDS_CODES else 'wdi'}_{code}.json"


def fetch(url, name):
    for attempt in range(3):
        try:
            req = Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urlopen(req, timeout=90) as response:
                payload = response.read()
                final_url = response.url
                content_type = response.headers.get('Content-Type')
            if name.endswith('.json'):
                parsed = json.loads(payload)
                advanced = isinstance(parsed, dict) and isinstance(parsed.get('source'), dict) and isinstance(parsed['source'].get('data'), list)
                standard = isinstance(parsed, list) and len(parsed) == 2 and isinstance(parsed[1], list)
                if not (advanced or standard):
                    raise ValueError(f'Respuesta inesperada para {name}: {str(parsed)[:150]}')
                header = parsed if advanced else parsed[0]
                assert int(header.get('pages', 1)) == 1, 'Respuesta paginada incompleta'
            if name.endswith('.xlsx'):
                assert payload.startswith(b'PK'), 'Archivo XLSX inválido'
            if name.endswith('.pdf'):
                assert payload.startswith(b'%PDF'), 'Archivo PDF inválido'
            if name.endswith('.csv'):
                assert b'gii_2023' in payload, 'Columnas GII ausentes'
            (RAW / name).write_bytes(payload)
            return dict(archivo=name, url=url, url_final=final_url, tipo_contenido=content_type,
                        descarga_utc=datetime.now(timezone.utc).isoformat(), bytes=len(payload),
                        sha256=hashlib.sha256(payload).hexdigest())
        except Exception:
            if attempt == 2:
                raise
            sleep(2 * (attempt + 1))


def download_wdi(spec):
    _, code, _, _ = spec
    base = 'https://api.worldbank.org/v2'
    countries = ';'.join(COUNTRIES)
    source = 6 if code in IDS_CODES else 2
    if code in IDS_CODES:
        url = f'{base}/sources/6/country/IND/counterpart-area/WLD/series/{code}/data?format=json&per_page=20000'
    else:
        url = f'{base}/country/{countries}/indicator/{code}?source=2&format=json&per_page=20000'
    return [fetch(url, series_file(code)),
            fetch(f'{base}/indicator/{code}?source={source}&format=json&per_page=20000', f'metadata_{code}.json')]


def save_csv(name, rows, fields=None):
    if fields is None:
        fields = list(rows[0])
    path = OUT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter=';')
        writer.writeheader()
        for row in rows:
            writer.writerow({k: str(v).replace('.', ',') if isinstance(v, float) else v for k, v in row.items()})


def read_json(name):
    return json.loads((RAW / name).read_bytes())


def process_wdi():
    long, dictionary, coverage, indexed, excluded = [], [], [], {}, []
    for alias, code, unit, block in SERIES:
        data = read_json(series_file(code))
        source = 6 if code in IDS_CODES else 2
        if code in IDS_CODES:
            records = []
            for raw_row in data['source']['data']:
                dims = {v['concept'].lower(): v for v in raw_row['variable']}
                assert dims['counterpart-area']['id'] == 'WLD'
                year = int(dims['time']['value'])
                if year > IDS_HISTORICAL_END:
                    if raw_row['value'] is not None:
                        excluded.append(dict(pais=dims['country']['id'], anio=year, variable=alias,
                                             codigo=code, valor=raw_row['value'], unidad=unit,
                                             estado='fuera_del_periodo_historico_IDS_no_usar_como_observado',
                                             archivo_raw=series_file(code)))
                    continue
                records.append(dict(countryiso3code=dims['country']['id'], date=dims['time']['value'],
                                    indicator={'id': dims['series']['id']}, value=raw_row['value']))
            data = [dict(lastupdated=data.get('lastupdated')), records]
        meta = read_json(f'metadata_{code}.json')[1][0]
        assert meta['id'] == code
        dictionary.append(dict(variable=alias, codigo=code, punto=1 if block == 'demografia' else 2,
                               bloque=block, unidad=unit, fuente_id=source,
                               nombre_original=meta['name'], definicion=meta.get('sourceNote', ''),
                               fuente=meta.get('sourceOrganization', ''),
                               url=f'https://api.worldbank.org/v2/indicator/{code}?source={source}&format=json'))
        for row in data[1]:
            country, year = row['countryiso3code'], int(row['date'])
            assert country in COUNTRIES and row['indicator']['id'] == code
            key = (country, year, alias)
            assert key not in indexed, f'Duplicado {key}'
            value = row['value']
            assert value is None or (isinstance(value, (int, float)) and math.isfinite(value))
            indexed[key] = value
            long.append(dict(pais=country, anio=year, variable=alias, codigo=code, valor=value,
                             unidad=unit, estado_observacion=row.get('obs_status', ''),
                             actualizacion_fuente=data[0].get('lastupdated'), archivo_raw=series_file(code)))
        for country in COUNTRIES:
            years = sorted(y for c, y, a in indexed if c == country and a == alias and indexed[c, y, a] is not None)
            coverage.append(dict(pais=country, variable=alias, codigo=code, observaciones=len(years),
                                 primer_anio=years[0] if years else None, ultimo_anio=years[-1] if years else None,
                                 anios_disponibles='|'.join(map(str, years)),
                                 actualizacion_fuente=data[0].get('lastupdated'),
                                 estado='fuera_de_cobertura_IDS' if source == 6 and country == 'USA' else ('disponible' if years else 'sin_datos')))
    natural_alias = 'crecimiento_natural_pct_anual'
    natural_code = 'CALC.CREC.NAT'
    reference_update = read_json('wdi_SP.DYN.CBRT.IN.json')[0].get('lastupdated')
    dictionary.append(dict(variable=natural_alias, codigo=natural_code, punto=1, bloque='demografia',
                           unidad='% anual', fuente_id='', nombre_original='Crecimiento natural de la población',
                           definicion='(natalidad por mil - mortalidad por mil) / 10; sin componente migratorio',
                           fuente='Elaboración propia con WDI SP.DYN.CBRT.IN y SP.DYN.CDRT.IN', url=''))
    for country in COUNTRIES:
        available = []
        for year in sorted({y for c, y, a in indexed if c == country}):
            birth = indexed.get((country, year, 'natalidad_por_mil'))
            death = indexed.get((country, year, 'mortalidad_por_mil'))
            value = (birth - death) / 10 if birth is not None and death is not None else None
            indexed[country, year, natural_alias] = value
            if value is not None:
                available.append(year)
            long.append(dict(pais=country, anio=year, variable=natural_alias, codigo=natural_code, valor=value,
                             unidad='% anual', estado_observacion='calculado' if value is not None else '',
                             actualizacion_fuente=reference_update,
                             archivo_raw='wdi_SP.DYN.CBRT.IN.json | wdi_SP.DYN.CDRT.IN.json'))
        coverage.append(dict(pais=country, variable=natural_alias, codigo=natural_code,
                             observaciones=len(available), primer_anio=min(available) if available else None,
                             ultimo_anio=max(available) if available else None,
                             anios_disponibles='|'.join(map(str, available)),
                             actualizacion_fuente=reference_update, estado='calculado'))
    valid_years = [y for (c, y, a), v in indexed.items() if v is not None]
    panel = [dict(pais=c, nombre_pais=name, anio=y,
                  **{s[0]: indexed.get((c, y, s[0])) for s in SERIES},
                  **{natural_alias: indexed.get((c, y, natural_alias))})
             for c, name in COUNTRIES.items() for y in range(min(valid_years), max(valid_years) + 1)]
    save_csv('apoyo/panel_anual.csv', panel)
    save_csv('apoyo/series_largas.csv', sorted(long, key=lambda r: (r['pais'], r['variable'], r['anio'])))
    save_csv('apoyo/diccionario_variables.csv', dictionary)
    save_csv('apoyo/cobertura_fuentes.csv', coverage)
    save_csv('apoyo/ids_valores_fuera_periodo_historico.csv', excluded,
             ['pais', 'anio', 'variable', 'codigo', 'valor', 'unidad', 'estado', 'archivo_raw'])
    # Bases principales por punto; crecimiento natural calculado con tasas brutas.
    for name, blocks in [('main_punto_1', {'demografia'}), ('main_punto_2', {'migracion', 'divisas'})]:
        fields = ['pais', 'nombre_pais', 'anio'] + [s[0] for s in SERIES if s[3] in blocks]
        if name == 'main_punto_1':
            fields.insert(fields.index('crecimiento_poblacional_pct_anual') + 1, natural_alias)
        save_csv(f'{name}.csv', [{k: r[k] for k in fields} for r in panel])
    return coverage


GII_VARS = {
    'gii': ('gii', 'índice 0-1; menor valor indica menor desigualdad', 'indice'),
    'mmr': ('mortalidad_materna', 'muertes maternas por 100.000 nacidos vivos', 'salud_reproductiva'),
    'abr': ('natalidad_adolescente', 'nacimientos por 1.000 mujeres de 15-19 años', 'salud_reproductiva'),
    'pr_f': ('bancas_parlamento_mujeres', '% de bancas', 'empoderamiento'),
    'pr_m': ('bancas_parlamento_varones', '% de bancas', 'empoderamiento'),
    'se_f': ('secundaria_mujeres', '% de mujeres de 25+ con al menos algo de secundaria', 'empoderamiento'),
    'se_m': ('secundaria_varones', '% de varones de 25+ con al menos algo de secundaria', 'empoderamiento'),
    'lfpr_f': ('participacion_laboral_mujeres', '% de mujeres de 15+', 'mercado_laboral'),
    'lfpr_m': ('participacion_laboral_varones', '% de varones de 15+', 'mercado_laboral'),
}


def process_gii():
    from openpyxl import load_workbook
    payload = (RAW / HDR_FILES[0]).read_bytes()
    try:
        text = payload.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = payload.decode('cp1252')
    global_rows = list(csv.DictReader(io.StringIO(text)))
    countries = [r for r in global_rows if r['iso3'] in COUNTRIES]
    assert len(countries) == len(COUNTRIES)
    panel, long, coverage = [], [], []
    for row in countries:
        for year in range(1990, 2024):
            record = dict(pais=row['iso3'], nombre_pais=COUNTRIES[row['iso3']], anio=year)
            for prefix, (alias, unit, dimension) in GII_VARS.items():
                raw = row.get(f'{prefix}_{year}', '').strip()
                value = float(raw) if raw not in ('', '..') else None
                record[alias] = value
                long.append(dict(pais=row['iso3'], anio=year, variable=alias, codigo_pnud=prefix,
                                 valor=value, unidad=unit, dimension=dimension,
                                 anio_referencia_indicador=min(year, 2020) if prefix == 'mmr' else year,
                                 archivo_raw=HDR_FILES[0]))
            panel.append(record)
        for prefix, (alias, unit, dimension) in GII_VARS.items():
            available = [r['anio'] for r in panel if r['pais'] == row['iso3'] and r[alias] is not None]
            coverage.append(dict(pais=row['iso3'], variable=alias, observaciones=len(available),
                                 primer_anio=min(available) if available else None,
                                 ultimo_anio=max(available) if available else None))
    save_csv('genero/gii_componentes_panel.csv', panel)
    save_csv('genero/gii_componentes_largo.csv', long)
    save_csv('genero/gii_evolucion.csv', [{k: r[k] for k in ['pais', 'nombre_pais', 'anio', 'gii']} for r in panel])
    save_csv('genero/cobertura_gii.csv', coverage)
    metadata_wb = load_workbook(RAW / HDR_FILES[1], read_only=True, data_only=True)
    codebook = {str(row[1]): {'definicion_original': row[0], 'periodo_fuente': row[2]}
                for row in metadata_wb['codebook'].iter_rows(values_only=True) if row[1] is not None}
    save_csv('genero/diccionario_gii.csv', [dict(variable=v[0], codigo_pnud=k, unidad=v[1], dimension=v[2],
                                        **codebook.get(k, {})) for k, v in GII_VARS.items()])
    metadata_wb.close()
    # Se extrae además la tabla publicada: preserva periodos y notas originales.
    wb = load_workbook(RAW / HDR_FILES[2], read_only=True, data_only=True)
    dump = {sheet.title: [list(r) for r in sheet.iter_rows(values_only=True)] for sheet in wb}
    (OUT / 'genero/gii_tabla_publicada_celdas.json').write_text(json.dumps(dump, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    table = dump['Table 5. GII']
    notes = {str(r[1]).split('. ', 1)[0]: str(r[1]).split('. ', 1)[1] for r in table
             if isinstance(r[1], str) and len(r[1]) > 3 and r[1][0].islower() and r[1][1:3] == '. '}
    columns = [(2, 'gii', 'índice 0-1', 'indice'), (4, 'ranking_gii', 'puesto', 'indice'),
               (6, 'mortalidad_materna', GII_VARS['mmr'][1], 'salud_reproductiva'),
               (8, 'natalidad_adolescente', GII_VARS['abr'][1], 'salud_reproductiva'),
               (10, 'bancas_parlamento_mujeres', '% de bancas', 'empoderamiento'),
               (12, 'secundaria_mujeres', GII_VARS['se_f'][1], 'empoderamiento'),
               (14, 'secundaria_varones', GII_VARS['se_m'][1], 'empoderamiento'),
               (16, 'participacion_laboral_mujeres', GII_VARS['lfpr_f'][1], 'mercado_laboral'),
               (18, 'participacion_laboral_varones', GII_VARS['lfpr_m'][1], 'mercado_laboral')]
    comparison = []
    lookup = {'India': 'IND', 'United States': 'USA'}
    for row_number, row in enumerate(table, 1):
        if row[1] not in lookup:
            continue
        for col, variable, unit, dimension in columns:
            marks = [str(value) for value in (table[6][col + 1], row[col + 1]) if value]
            if col >= 16:
                marks.append('a')
            comparison.append(dict(pais=lookup[row[1]], variable=variable, dimension=dimension,
                                    valor=row[col], unidad=unit, anio_encabezado=table[6][col],
                                    notas=' | '.join(notes.get(mark, mark) for mark in dict.fromkeys(marks)),
                                    fila_excel=row_number, columna_excel=col + 1, archivo_raw=HDR_FILES[2]))
    assert len(comparison) == len(COUNTRIES) * len(columns)
    save_csv('genero/gii_tabla5_comparacion.csv', comparison)
    wb.close()
    return coverage


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--offline', action='store_true', help='Regenerar desde las fuentes ya descargadas')
    args = parser.parse_args()
    RAW.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    if not args.offline:
        with ThreadPoolExecutor(max_workers=4) as pool:
            entries = [item for group in pool.map(download_wdi, SERIES) for item in group]
        # Guardar antes de descargar PNUD para conservar trazabilidad ante un fallo.
        (RAW / 'manifest.json').write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding='utf-8')
        for filename in HDR_FILES:
            entries.append(fetch(HDR_BASE + filename, filename))
            (RAW / 'manifest.json').write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding='utf-8')
    for entry in read_json('manifest.json'):
        assert hashlib.sha256((RAW / entry['archivo']).read_bytes()).hexdigest() == entry['sha256']
    coverage = process_wdi()
    gii_coverage = process_gii()
    for row in coverage:
        print(f"{row['pais']} {row['variable']}: {row['observaciones']} ({row['primer_anio']}-{row['ultimo_anio']})")
    print('GII:', [r for r in gii_coverage if r['variable'] == 'gii'])


if __name__ == '__main__':
    main()
