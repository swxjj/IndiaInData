import os
import json
import urllib.request
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# -----------------------------------------------------------------------------
# CONFIGURACIÓN DE PARÁMETROS Y RUTAS
# -----------------------------------------------------------------------------
COUNTRY = "IND"
COUNTRY_NAME = "India"

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_DIR = os.path.join(BASE_DIR, "data", "raw")
PROC_DIR = os.path.join(BASE_DIR, "data", "processed")
FIG_DIR = os.path.join(BASE_DIR, "output", "figures")
OUT_DIR = os.path.join(BASE_DIR, "output")

os.makedirs(RAW_DIR, exist_ok=True)
os.makedirs(PROC_DIR, exist_ok=True)
os.makedirs(FIG_DIR, exist_ok=True)

INDICATORS = {
    "NV.AGR.TOTL.KD.ZG": "agriculture_va_growth_pct",   # VA Agricultura (% crec. anual)
    "NV.IND.TOTL.KD.ZG": "industry_va_growth_pct",      # VA Industria (% crec. anual)
    "NV.SRV.TOTL.KD.ZG": "services_va_growth_pct",      # VA Servicios (% crec. anual)
    "NY.GDP.PCAP.KD.ZG": "gdp_pcap_growth_annual_pct"   # PIB per cápita (% crec. anual)
}

# -----------------------------------------------------------------------------
# 1. DESCARGA DE DATOS SECTORIALES DESDE WORLD BANK API
# -----------------------------------------------------------------------------
def fetch_sector_data():
    print("=" * 80)
    print(" TP3 PARTE 2: DESCARGA DE DATOS DE CRECIMIENTO SECTORIAL (INDIA)")
    print("=" * 80)
    
    raw_records = []
    dict_by_ind = {}
    
    for code, col in INDICATORS.items():
        url = f"https://api.worldbank.org/v2/country/{COUNTRY}/indicator/{code}?format=json&per_page=3000"
        print(f"-> Descargando {code:<22} ({col})...", end=" ", flush=True)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            series = {}
            if len(data) > 1 and data[1]:
                for row in data[1]:
                    if row.get('value') is not None:
                        yr = int(row['date'])
                        val = float(row['value'])
                        series[yr] = val
                        raw_records.append({
                            "country_iso": COUNTRY,
                            "country_name": COUNTRY_NAME,
                            "indicator_code": code,
                            "indicator_name": col,
                            "year": yr,
                            "value": val
                        })
            print(f"OK: {len(series)} registros ({min(series.keys())}-{max(series.keys())})")
            dict_by_ind[col] = series
            
    # Guardar raw CSV
    df_raw = pd.DataFrame(raw_records)
    raw_path = os.path.join(RAW_DIR, "wb_india_sector_growth_raw.csv")
    df_raw.to_csv(raw_path, index=False, encoding="utf-8")
    print(f"   [Guardado Raw]: {raw_path}")
    
    # Construir dataframe procesado
    df_sec = pd.DataFrame(dict_by_ind).sort_index().reset_index().rename(columns={'index': 'year'})
    df_sec['country_iso'] = COUNTRY
    df_sec['country_name'] = COUNTRY_NAME
    
    # Reordenar columnas base
    cols_base = ['country_iso', 'country_name', 'year', 'agriculture_va_growth_pct', 'industry_va_growth_pct', 'services_va_growth_pct', 'gdp_pcap_growth_annual_pct']
    df_sec = df_sec[cols_base].dropna().reset_index(drop=True)
    
    # -------------------------------------------------------------------------
    # 2. CÁLCULO DE DESVIACIÓN ESTÁNDAR SECTORIAL ANUAL (NURKSE VS HIRSCHMAN)
    # -------------------------------------------------------------------------
    sector_cols = ['agriculture_va_growth_pct', 'industry_va_growth_pct', 'services_va_growth_pct']
    df_sec['sector_growth_std'] = df_sec[sector_cols].std(axis=1)
    
    # Sector líder
    def get_leading_sector(row):
        m = max(row['agriculture_va_growth_pct'], row['industry_va_growth_pct'], row['services_va_growth_pct'])
        if m == row['services_va_growth_pct']:
            return "Servicios"
        elif m == row['industry_va_growth_pct']:
            return "Industria"
        else:
            return "Agricultura"
            
    df_sec['leading_sector'] = df_sec.apply(get_leading_sector, axis=1)
    
    # Guardar CSV procesado
    sec_csv_path = os.path.join(PROC_DIR, "tp3_sector_growth_ind.csv")
    df_sec.to_csv(sec_csv_path, index=False, encoding="utf-8")
    print(f"   [Guardado CSV Procesado]: {sec_csv_path} ({len(df_sec)} anios: {df_sec['year'].min()}-{df_sec['year'].max()})")
    
    return df_sec

# -----------------------------------------------------------------------------
# 3. ACTUALIZAR LIBRO EXCEL CON HOJA DE CRECIMIENTO SECTORIAL
# -----------------------------------------------------------------------------
def update_excel_workbook(df_sec):
    print("\n--- ACTUALIZANDO LIBRO EXCEL CON HOJA DE CRECIMIENTO SECTORIAL ---")
    excel_paths = [
        os.path.join(PROC_DIR, "india_tp3_datos.xlsx"),
        os.path.join(OUT_DIR, "india_tp3_datos.xlsx")
    ]
    
    for path in excel_paths:
        if not os.path.exists(path):
            continue
            
        # Leer hojas existentes
        with pd.ExcelFile(path) as xls:
            sheets_dict = {sheet: pd.read_excel(xls, sheet_name=sheet) for sheet in xls.sheet_names}
            
        # Agregar / sobrescribir hoja Crecimiento_Sectorial
        sheets_dict["Crecimiento_Sectorial"] = df_sec
        
        # Escribir con formato
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for sheet_name, data in sheets_dict.items():
                data.to_excel(writer, sheet_name=sheet_name, index=False)
                
            wb = writer.book
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            fill_header = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid") # Púrpura elegante para crecimiento sectorial
            
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            
            ws = wb["Crecimiento_Sectorial"]
            ws.views.sheetView[0].showGridLines = True
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = fill_header
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
            ws.row_dimensions[1].height = 28
            
            for col in ws.columns:
                col_letter = col[0].column_letter
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
                for cell in col[1:]:
                    cell.border = thin_border
                    if isinstance(cell.value, float):
                        cell.number_format = "#,##0.00"
                    elif isinstance(cell.value, int) and col_letter == 'C':
                        cell.number_format = "0"
                        
        print(f"   [Excel actualizado con hoja Crecimiento_Sectorial]: {path}")

# -----------------------------------------------------------------------------
# 4. GRAFICACIÓN ACADÉMICA LIMPIA (HISTÓRICO COMPLETO Y FOCALIZADO DESDE 1990)
# -----------------------------------------------------------------------------
def plot_growth_charts(df_sec):
    print("\n--- GENERACIÓN DE GRÁFICOS ACADÉMICOS DE CRECIMIENTO SECTORIAL ---")
    
    plt.style.use('seaborn-v0_8-whitegrid' if 'seaborn-v0_8-whitegrid' in plt.style.available else 'default')
    plt.rcParams['font.family'] = 'DejaVu Sans'
    plt.rcParams['font.size'] = 10
    plt.rcParams['axes.titlesize'] = 11
    plt.rcParams['axes.labelsize'] = 10
    plt.rcParams['legend.fontsize'] = 9
    plt.rcParams['figure.titlesize'] = 13
    
    # -------------------------------------------------------------------------
    # FIGURA 5: Histórico Completo de Crecimiento Sectorial (1961–2025)
    # -------------------------------------------------------------------------
    print("-> Generando Gráfico 5: Crecimiento Sectorial Histórico Completo (1961–2025)...")
    fig5, (ax5a, ax5b) = plt.subplots(2, 1, figsize=(11.5, 8.5), sharex=True, gridspec_kw={'height_ratios': [1.6, 1]})
    
    # Panel superior: Crecimiento por sector vs PIB pc
    ax5a.plot(df_sec['year'], df_sec['agriculture_va_growth_pct'], label='Agricultura, silvicultura y pesca', color='#2ca02c', lw=1.8)
    ax5a.plot(df_sec['year'], df_sec['industry_va_growth_pct'], label='Industria y construcción', color='#d62728', lw=1.8)
    ax5a.plot(df_sec['year'], df_sec['services_va_growth_pct'], label='Servicios', color='#1f77b4', lw=2.0)
    ax5a.plot(df_sec['year'], df_sec['gdp_pcap_growth_annual_pct'], label='PIB per cápita (% crec. anual)', color='#222222', lw=2.2, ls='--', alpha=0.9)
    
    ax5a.axhline(0, color='gray', linestyle='-', lw=0.9)
    ax5a.set_ylabel('Tasa de Crecimiento Anual (%)')
    ax5a.set_title("Panel A: Tasa de Crecimiento Anual por Sector vs. PIB per Cápita (1961–2025)", fontweight='bold')
    ax5a.legend(loc='lower left', frameon=True, ncol=2)
    ax5a.grid(True, linestyle='--', alpha=0.5)
    
    # Panel inferior: Desviación estándar transversal anual (Nurkse vs Hirschman)
    mean_std_hist = df_sec['sector_growth_std'].mean()
    ax5b.bar(df_sec['year'], df_sec['sector_growth_std'], color='#708090', alpha=0.6, width=0.7, label='Desviación estándar sectorial anual (DE)')
    ax5b.axhline(mean_std_hist, color='#b22222', linestyle='--', lw=1.8, label=f"Media histórica DE ({mean_std_hist:.2f} p.p.)")
    
    ax5b.set_ylabel('Desviación Estándar (p.p.)')
    ax5b.set_xlabel('Año')
    ax5b.set_xlim(1960, 2026)
    ax5b.xaxis.set_major_locator(ticker.MultipleLocator(5))
    ax5b.set_title("Panel B: Dispersión Sectorial Anual – Grado de Desequilibrio (Nurkse vs. Hirschman)", fontweight='bold')
    ax5b.legend(loc='upper right', frameon=True)
    ax5b.grid(True, linestyle='--', alpha=0.5)
    
    fig5.suptitle("India: Crecimiento Equilibrado vs. Desequilibrado por Sectores (1961–2025)", fontweight='bold', y=0.98)
    fig5.tight_layout()
    fig5_path = os.path.join(FIG_DIR, "fig5_india_crecimiento_sectorial_historico.png")
    fig5.savefig(fig5_path, dpi=300)
    plt.close(fig5)
    print(f"   [OK]: {fig5_path}")
    
    # -------------------------------------------------------------------------
    # FIGURA 6: Focalizado desde 1990 (según consigna oficial del TP3)
    # -------------------------------------------------------------------------
    print("-> Generando Gráfico 6: Crecimiento Sectorial desde 1990 (Focalizado TP3)...")
    df_post90 = df_sec[df_sec['year'] >= 1990].copy().reset_index(drop=True)
    
    fig6, (ax6a, ax6b) = plt.subplots(2, 1, figsize=(11.5, 8.5), sharex=True, gridspec_kw={'height_ratios': [1.6, 1]})
    
    # Panel superior: Sectores vs PIB pc desde 1990
    ax6a.plot(df_post90['year'], df_post90['agriculture_va_growth_pct'], label='Agricultura, silvicultura y pesca', color='#2ca02c', lw=2.0, marker='o', ms=3.5)
    ax6a.plot(df_post90['year'], df_post90['industry_va_growth_pct'], label='Industria y construcción', color='#d62728', lw=2.0, marker='s', ms=3.5)
    ax6a.plot(df_post90['year'], df_post90['services_va_growth_pct'], label='Servicios', color='#1f77b4', lw=2.2, marker='^', ms=4.0)
    ax6a.plot(df_post90['year'], df_post90['gdp_pcap_growth_annual_pct'], label='PIB per cápita (% crec. anual)', color='#111111', lw=2.2, ls='--', alpha=0.85)
    
    ax6a.axhline(0, color='gray', linestyle='-', lw=0.9)
    ax6a.set_ylabel('Tasa de Crecimiento Anual (%)')
    ax6a.set_title("Panel A: Tasa de Crecimiento Anual por Sector vs. PIB per Cápita (1990–2025)", fontweight='bold')
    ax6a.legend(loc='lower left', frameon=True, ncol=2)
    ax6a.grid(True, linestyle='--', alpha=0.5)
    
    # Panel inferior: Desviación estándar transversal desde 1990
    mean_std_90 = df_post90['sector_growth_std'].mean()
    ax6b.bar(df_post90['year'], df_post90['sector_growth_std'], color='#4682b4', alpha=0.6, width=0.65, label='Desviación estándar sectorial anual (DE)')
    ax6b.axhline(mean_std_90, color='#b22222', linestyle='--', lw=1.8, label=f"Promedio período 1990-2025 ({mean_std_90:.2f} p.p.)")
    
    ax6b.set_ylabel('Desviación Estándar (p.p.)')
    ax6b.set_xlabel('Año')
    ax6b.set_xlim(1989.5, 2025.5)
    ax6b.xaxis.set_major_locator(ticker.MultipleLocator(2))
    ax6b.set_title("Panel B: Desviación Estándar entre Crecimientos Sectoriales (Dispersión Anual)", fontweight='bold')
    ax6b.legend(loc='upper right', frameon=True)
    ax6b.grid(True, linestyle='--', alpha=0.5)
    
    fig6.suptitle("India: Dinámica de Crecimiento Sectorial y Desequilibrio Empírico (1990–2025)", fontweight='bold', y=0.98)
    fig6.tight_layout()
    fig6_path = os.path.join(FIG_DIR, "fig6_india_crecimiento_sectorial_desde_1990.png")
    fig6.savefig(fig6_path, dpi=300)
    plt.close(fig6)
    print(f"   [OK]: {fig6_path}")

def main():
    df_sec = fetch_sector_data()
    update_excel_workbook(df_sec)
    plot_growth_charts(df_sec)
    print("\n" + "=" * 80)
    print(" ¡PROCESO DE CRECIMIENTO SECTORIAL FINALIZADO CON ÉXITO!")
    print("=" * 80)

if __name__ == "__main__":
    main()
