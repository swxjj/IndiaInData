import os
import re
import csv
import json
import time
import urllib.request
import urllib.error
import pandas as pd

# -----------------------------------------------------------------------------
# CONFIGURACIÓN Y PARÁMETROS
# -----------------------------------------------------------------------------
COUNTRY = "IND"
COUNTRY_NAME = "India"
ECONSTATS_URL = "https://www.econstats.com/wdi/wdic_IND.htm"

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_DIR = os.path.join(BASE_DIR, "data", "raw")
PROCESSED_DIR = os.path.join(BASE_DIR, "data", "processed")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

os.makedirs(RAW_DIR, exist_ok=True)
os.makedirs(PROCESSED_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

# -----------------------------------------------------------------------------
# DICCIONARIO DE INDICADORES WORLD BANK (API OFICIAL)
# -----------------------------------------------------------------------------
WB_INDICATORS_GROWTH_INV = {
    "NY.GDP.PCAP.KD.ZG": {
        "col": "gdp_pcap_growth_pct",
        "name_es": "Crecimiento anual del PIB per cápita (%)",
        "name_en": "GDP per capita growth (annual %)",
        "unit": "%",
        "part": "Parte 1: Crecimiento e Inversión"
    },
    "NE.GDI.TOTL.ZS": {
        "col": "gross_capital_formation_pct_gdp",
        "name_es": "Inversión bruta (% del PIB)",
        "name_en": "Gross capital formation (% of GDP)",
        "unit": "% del PIB",
        "part": "Parte 1: Crecimiento e Inversión"
    },
    "NE.GDI.FTOT.ZS": {
        "col": "gross_fixed_capital_formation_pct_gdp",
        "name_es": "Inversión fija bruta (% del PIB)",
        "name_en": "Gross fixed capital formation (% of GDP)",
        "unit": "% del PIB",
        "part": "Parte 1: Crecimiento e Inversión"
    },
    "BX.KLT.DINV.WD.GD.ZS": {
        "col": "fdi_net_inflows_pct_gdp",
        "name_es": "IED neta (% del PIB)",
        "name_en": "Foreign direct investment, net inflows (% of GDP)",
        "unit": "% del PIB",
        "part": "Parte 1: Crecimiento e Inversión"
    }
}

WB_INDICATORS_INFRA = {
    "EG.ELC.ACCS.ZS": {
        "col": "wb_access_electricity_pct_pop",
        "name_es": "Acceso a la electricidad (% población)",
        "name_en": "Access to electricity (% of population)",
        "unit": "% población",
        "part": "Parte 2: Infraestructura (Servicios Básicos)"
    },
    "SH.H2O.BASW.ZS": {
        "col": "wb_basic_drinking_water_pct_pop",
        "name_es": "Acceso a agua potable básica (% población)",
        "name_en": "People using at least basic drinking water services (% of pop)",
        "unit": "% población",
        "part": "Parte 2: Infraestructura (Servicios Básicos)"
    },
    "IT.NET.USER.ZS": {
        "col": "wb_internet_users_pct_pop",
        "name_es": "Usuarios de Internet (% población)",
        "name_en": "Individuals using the Internet (% of population)",
        "unit": "% población",
        "part": "Parte 2: Infraestructura (Comunicaciones)"
    },
    "IT.CEL.SETS.P2": {
        "col": "wb_mobile_subscriptions_per_100",
        "name_es": "Telefonía celular (suscripciones por 100 hab)",
        "name_en": "Mobile cellular subscriptions (per 100 people)",
        "unit": "Por 100 hab.",
        "part": "Parte 2: Infraestructura (Comunicaciones)"
    },
    "IT.NET.BBND.P2": {
        "col": "wb_broadband_subscriptions_per_100",
        "name_es": "Banda ancha fija (suscripciones por 100 hab)",
        "name_en": "Fixed broadband subscriptions (per 100 people)",
        "unit": "Por 100 hab.",
        "part": "Parte 2: Infraestructura (Comunicaciones)"
    },
    "IS.AIR.PSGR": {
        "col": "wb_air_passengers_carried",
        "name_es": "Transporte aéreo de pasajeros",
        "name_en": "Air transport, passengers carried",
        "unit": "Pasajeros",
        "part": "Parte 2: Infraestructura (Transporte)"
    },
    "EG.USE.ELEC.KH.PC": {
        "col": "wb_electric_power_kwh_pc",
        "name_es": "Consumo eléctrico per cápita (kWh)",
        "name_en": "Electric power consumption (kWh per capita)",
        "unit": "kWh per cápita",
        "part": "Parte 2: Infraestructura (Energía)"
    }
}

# -----------------------------------------------------------------------------
# FUNCIONES DE DESCARGA WORLD BANK API
# -----------------------------------------------------------------------------
def fetch_wb_indicator(indicator_code, retries=3):
    url = f"https://api.worldbank.org/v2/country/{COUNTRY}/indicator/{indicator_code}?format=json&per_page=3000"
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
    
    for attempt in range(1, retries + 1):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=35) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                records = []
                if len(data) > 1 and data[1]:
                    for row in data[1]:
                        if row.get("value") is not None:
                            records.append({
                                "country_iso": COUNTRY,
                                "country_name": COUNTRY_NAME,
                                "year": int(row["date"]),
                                "indicator_id": indicator_code,
                                "value": float(row["value"])
                            })
                return records
        except Exception as e:
            print(f"   [Reintento {attempt}/{retries}] Error al descargar {indicator_code}: {e}")
            time.sleep(2)
    return []

# -----------------------------------------------------------------------------
# FUNCIONES DE DESCARGA Y PARSEO ECONSTATS
# -----------------------------------------------------------------------------
def fetch_econstats_india():
    print(f"\n-> Descargando tabla histórica completa de EconStats para India ({ECONSTATS_URL})...")
    req = urllib.request.Request(
        ECONSTATS_URL,
        headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:120.0) Gecko/20100101 Firefox/120.0"}
    )
    with urllib.request.urlopen(req, timeout=45) as resp:
        html = resp.read().decode("latin-1", errors="replace")
    print(f"   Descarga exitosa: {len(html):,} bytes.")
    return html

def parse_econstats_india(html):
    print("-> Parseando estructura de EconStats India...")
    tables = re.findall(r"<table[^>]*>.*?</table>", html, re.DOTALL | re.I)
    target_table = None
    for t in tables:
        if len(t) > 50000 and "IND" in t and "1960" in t:
            target_table = t
            break
            
    if not target_table:
        raise ValueError("No se pudo localizar la tabla principal en la página de EconStats.")
        
    rows = re.split(r"<tr[^>]*>", target_table, flags=re.I)
    header_cells = [re.sub(r"<[^<]+?>", "", c).strip() for c in re.split(r"<td[^>]*>", rows[1], flags=re.I)[1:]]
    year_cols = [(idx, int(c)) for idx, c in enumerate(header_cells) if c.isdigit() and len(c) == 4]
    
    print(f"   Detectadas {len(year_cols)} columnas anuales ({year_cols[0][1]} a {year_cols[-1][1]}).")
    
    all_records = []
    current_sec = "GENERAL"
    
    for r in rows[2:]:
        cells = [re.sub(r"<[^<]+?>", "", c).strip().replace("\xa0", " ").replace("&nbsp;", " ").strip() for c in re.split(r"<td[^>]*>", r, flags=re.I)[1:]]
        
        # Detectar fila de sección
        if len(cells) >= 4 and not cells[0].isdigit() and cells[3] and not cells[1]:
            current_sec = cells[3].strip()
            continue
            
        if len(cells) >= 5 and cells[1] == "IND":
            indicator_code = cells[3].strip()
            indicator_name = cells[4].strip()
            
            if not indicator_code:
                continue
                
            for y_idx, yr in year_cols:
                if y_idx < len(cells):
                    raw_val = cells[y_idx].replace(",", "").strip()
                    if raw_val and raw_val not in ("..", "-", "n.a.", ""):
                        try:
                            val = float(raw_val)
                            all_records.append({
                                "country_iso": "IND",
                                "country_name": "India",
                                "section": current_sec,
                                "indicator_code": indicator_code,
                                "indicator_name": indicator_name,
                                "year": yr,
                                "value": val
                            })
                        except ValueError:
                            pass
                            
    print(f"   Parseo completado: {len(all_records):,} observaciones extraídas para India.")
    return all_records

# -----------------------------------------------------------------------------
# EJECUCIÓN PRINCIPAL
# -----------------------------------------------------------------------------
def main():
    print("=" * 80)
    print(" TP3: DESCARGA Y ORGANIZACIÓN DE DATOS PARA INDIA (WORLD BANK Y ECONSTATS)")
    print("=" * 80)
    
    # 1. DESCARGA WORLD BANK (PARTE 1 Y PARTE 2)
    print("\n--- 1. DESCARGA DE INDICADORES DEL BANCO MUNDIAL (WDI API) ---")
    wb_records = []
    
    all_wb_indicators = {**WB_INDICATORS_GROWTH_INV, **WB_INDICATORS_INFRA}
    
    for code, meta in all_wb_indicators.items():
        print(f"-> Descargando {code:<22} | {meta['name_es'][:40]:<40}...", end=" ", flush=True)
        rec = fetch_wb_indicator(code)
        print(f"OK: {len(rec)} observaciones")
        wb_records.extend(rec)
        
    # Guardar raw WDI
    wb_raw_path = os.path.join(RAW_DIR, "wb_india_indicators_raw.csv")
    pd.DataFrame(wb_records).to_csv(wb_raw_path, index=False, encoding="utf-8")
    print(f"   [Guardado WDI Raw]: {wb_raw_path}")

    # 2. DESCARGA Y PARSEO ECONSTATS
    print("\n--- 2. DESCARGA Y PARSEO HISTÓRICO ECONSTATS ---")
    econstats_html = fetch_econstats_india()
    econstats_records = parse_econstats_india(econstats_html)
    df_econ = pd.DataFrame(econstats_records)
    
    # Guardar raw EconStats consolidado
    econ_all_raw_path = os.path.join(RAW_DIR, "econstats_india_all.csv")
    df_econ.to_csv(econ_all_raw_path, index=False, encoding="utf-8")
    print(f"   [Guardado EconStats Raw Completo]: {econ_all_raw_path}")
    
    # Filtrar subconjuntos raw de EconStats solicitados en consigna
    # Transporte (privilegiar)
    df_transport = df_econ[
        (df_econ['section'].str.contains('TRANSPORT', case=False, na=False)) |
        (df_econ['indicator_code'].str.startswith('IS.RO')) |
        (df_econ['indicator_code'].str.startswith('IS.RR')) |
        (df_econ['indicator_code'].str.startswith('IS.AIR')) |
        (df_econ['indicator_code'].str.startswith('IS.SHP')) |
        (df_econ['indicator_code'] == 'IE.PPI.TRAN.CD')
    ]
    df_transport.to_csv(os.path.join(RAW_DIR, "econstats_india_transport.csv"), index=False, encoding="utf-8")
    
    # Energía
    df_energy = df_econ[
        (df_econ['section'].str.contains('ENERGY', case=False, na=False)) |
        (df_econ['indicator_code'].str.startswith('EG.')) |
        (df_econ['indicator_code'] == 'IE.PPI.ENGY.CD')
    ]
    df_energy.to_csv(os.path.join(RAW_DIR, "econstats_india_energy.csv"), index=False, encoding="utf-8")
    
    # Comunicaciones
    df_comms = df_econ[
        (df_econ['section'].str.contains('COMMUNICATION', case=False, na=False)) |
        (df_econ['indicator_code'].str.startswith('IT.')) |
        (df_econ['indicator_code'] == 'IE.PPI.TELE.CD')
    ]
    df_comms.to_csv(os.path.join(RAW_DIR, "econstats_india_communications.csv"), index=False, encoding="utf-8")
    print(f"   [Guardados subconjuntos temáticos]: transport ({len(df_transport)} obs), energy ({len(df_energy)} obs), comms ({len(df_comms)} obs)")

    # -------------------------------------------------------------------------
    # 3. ORGANIZACIÓN Y ARMONIZACIÓN EN DOS PARTES
    # -------------------------------------------------------------------------
    print("\n--- 3. ARMONIZACIÓN Y PIVOT EN DOS PARTES TEMÁTICAS ---")
    
    # Años comunes de análisis: 1960 a 2025
    years_df = pd.DataFrame({"year": range(1960, 2026)})
    
    # --- PARTE 1: CRECIMIENTO E INVERSIÓN ---
    df_wb = pd.DataFrame(wb_records)
    part1_dfs = [years_df]
    
    for code, meta in WB_INDICATORS_GROWTH_INV.items():
        sub = df_wb[df_wb['indicator_id'] == code][['year', 'value']].rename(columns={'value': meta['col']})
        part1_dfs.append(sub)
        
    df_part1 = part1_dfs[0]
    for s in part1_dfs[1:]:
        df_part1 = pd.merge(df_part1, s, on='year', how='left')
        
    # Filtrar solo anios con al menos un dato
    df_part1 = df_part1.dropna(subset=[meta['col'] for meta in WB_INDICATORS_GROWTH_INV.values()], how='all')
    df_part1['country_iso'] = "IND"
    df_part1['country_name'] = "India"
    cols_order_p1 = ['country_iso', 'country_name', 'year'] + [meta['col'] for meta in WB_INDICATORS_GROWTH_INV.values()]
    df_part1 = df_part1[cols_order_p1].sort_values('year').reset_index(drop=True)
    
    # Guardar Parte 1 en CSV
    p1_csv_path = os.path.join(PROCESSED_DIR, "india_crecimiento_inversion.csv")
    df_part1.to_csv(p1_csv_path, index=False, encoding="utf-8")
    print(f"   [Parte 1 generada]: {p1_csv_path} ({len(df_part1)} anios: {df_part1['year'].min()}-{df_part1['year'].max()})")

    # --- PARTE 2: INFRAESTRUCTURA ---
    # Combinar indicadores de WDI + indicadores clave de EconStats
    infra_dict = {}
    
    # Agregar WDI infra
    for code, meta in WB_INDICATORS_INFRA.items():
        sub = df_wb[df_wb['indicator_id'] == code][['year', 'value']].set_index('year')['value']
        infra_dict[meta['col']] = sub
        
    # Agregar indicadores destacados de EconStats (Transporte, Energía, Telecomunicaciones)
    econstats_infra_vars = {
        # Transporte ferroviario (clave para Big Push en India)
        "IS.RRS.TOTL.KM": "econ_rail_lines_total_route_km",
        "IS.RRS.PASG.KM": "econ_railways_passengers_million_km",
        "IS.RRS.GOOD.MT.K6": "econ_railways_goods_million_ton_km",
        # Transporte carretero
        "IS.ROD.TOTL.KM": "econ_roads_total_network_km",
        "IS.ROD.PAVE.ZS": "econ_roads_paved_pct",
        # Transporte aéreo
        "IS.AIR.DPRT": "econ_air_carrier_departures",
        "IS.AIR.GOOD.MT.K1": "econ_air_freight_million_ton_km",
        # Inversión privada en transporte
        "IE.PPI.TRAN.CD": "econ_ppi_transport_usd",
        # Energía
        "EG.ELC.PROD.KH": "econ_electricity_production_kwh",
        "EG.ELC.COAL.ZS": "econ_electricity_from_coal_pct",
        "EG.ELC.HYRO.ZS": "econ_electricity_from_hydro_pct",
        "EG.USE.COMM.FO.ZS": "econ_fossil_fuel_energy_pct",
        "IE.PPI.ENGY.CD": "econ_ppi_energy_usd",
        # Telecomunicaciones
        "IT.MLT.MAIN.P2": "econ_telephone_lines_per_100",
        "IT.TEL.INVS.RV.ZS": "econ_telecom_investment_pct_revenue",
        "IE.PPI.TELE.CD": "econ_ppi_telecom_usd"
    }
    
    for code, col_name in econstats_infra_vars.items():
        sub = df_econ[df_econ['indicator_code'] == code][['year', 'value']].drop_duplicates(subset='year').set_index('year')['value']
        infra_dict[col_name] = sub
        
    df_part2 = pd.DataFrame(infra_dict, index=range(1960, 2026))
    df_part2.index.name = 'year'
    df_part2 = df_part2.reset_index()
    # Filtrar anios donde haya datos
    df_part2 = df_part2.dropna(how='all', subset=[c for c in df_part2.columns if c != 'year'])
    df_part2['country_iso'] = "IND"
    df_part2['country_name'] = "India"
    cols_order_p2 = ['country_iso', 'country_name', 'year'] + [c for c in df_part2.columns if c not in ('country_iso', 'country_name', 'year')]
    df_part2 = df_part2[cols_order_p2].sort_values('year').reset_index(drop=True)
    
    # Guardar Parte 2 en CSV
    p2_csv_path = os.path.join(PROCESSED_DIR, "india_infraestructura.csv")
    df_part2.to_csv(p2_csv_path, index=False, encoding="utf-8")
    print(f"   [Parte 2 generada]: {p2_csv_path} ({len(df_part2)} anios: {df_part2['year'].min()}-{df_part2['year'].max()}, {len(df_part2.columns)-3} variables)")

    # --- PANEL CONSOLIDADO COMPLETO ---
    df_panel = pd.merge(df_part1, df_part2, on=['country_iso', 'country_name', 'year'], how='outer').sort_values('year').reset_index(drop=True)
    panel_csv_path = os.path.join(PROCESSED_DIR, "tp3_panel_india.csv")
    df_panel.to_csv(panel_csv_path, index=False, encoding="utf-8")
    print(f"   [Panel Completo CSV]: {panel_csv_path}")

    # -------------------------------------------------------------------------
    # 4. EXPORTACIÓN A EXCEL CON MÚLTIPLES HOJAS Y FORMATO PROFESIONAL
    # -------------------------------------------------------------------------
    print("\n--- 4. EXPORTACIÓN A EXCEL MULTI-HOJA (.XLSX) ---")
    excel_path_proc = os.path.join(PROCESSED_DIR, "india_tp3_datos.xlsx")
    excel_path_out = os.path.join(OUTPUT_DIR, "india_tp3_datos.xlsx")
    
    # Metadatos para hoja Diccionario
    metadata_rows = []
    
    for code, meta in WB_INDICATORS_GROWTH_INV.items():
        col_name = meta['col']
        valid = df_part1.dropna(subset=[col_name])
        cob_str = f"{valid['year'].min()} - {valid['year'].max()} (n={len(valid)})" if not valid.empty else "Sin datos"
        metadata_rows.append({
            "Bloque": "Parte 1: Crecimiento e Inversión",
            "Columna": col_name,
            "Nombre Variable (Español)": meta['name_es'],
            "Nombre Original (Inglés)": meta['name_en'],
            "Código Indicador": code,
            "Fuente Oficial": "World Bank - World Development Indicators (WDI)",
            "Unidad de Medida": meta['unit'],
            "Cobertura": cob_str
        })
        
    for code, meta in WB_INDICATORS_INFRA.items():
        col_name = meta['col']
        valid = df_part2.dropna(subset=[col_name])
        cob_str = f"{valid['year'].min()} - {valid['year'].max()} (n={len(valid)})" if not valid.empty else "Sin datos"
        metadata_rows.append({
            "Bloque": meta['part'],
            "Columna": col_name,
            "Nombre Variable (Español)": meta['name_es'],
            "Nombre Original (Inglés)": meta['name_en'],
            "Código Indicador": code,
            "Fuente Oficial": "World Bank - World Development Indicators (WDI)",
            "Unidad de Medida": meta['unit'],
            "Cobertura": cob_str
        })
        
    econ_meta_desc = {
        "IS.RRS.TOTL.KM": ("Líneas de ferrocarril total (km de vía)", "Rail lines (total route-km)", "km", "Parte 2: Infraestructura (Transporte)"),
        "IS.RRS.PASG.KM": ("Pasajeros de ferrocarril (millones de pasajero-km)", "Railways, passengers carried (million passenger-km)", "Millones pas-km", "Parte 2: Infraestructura (Transporte)"),
        "IS.RRS.GOOD.MT.K6": ("Carga ferroviaria (millones de ton-km)", "Railways, goods transported (million ton-km)", "Millones ton-km", "Parte 2: Infraestructura (Transporte)"),
        "IS.ROD.TOTL.KM": ("Red vial total (km)", "Roads, total network (km)", "km", "Parte 2: Infraestructura (Transporte)"),
        "IS.ROD.PAVE.ZS": ("Carreteras pavimentadas (% de la red vial)", "Roads, paved (% of total roads)", "%", "Parte 2: Infraestructura (Transporte)"),
        "IS.AIR.DPRT": ("Despegues aéreos de transportistas registrados", "Air transport, registered carrier departures worldwide", "Despegues", "Parte 2: Infraestructura (Transporte)"),
        "IS.AIR.GOOD.MT.K1": ("Carga aérea (millones de ton-km)", "Air transport, freight (million ton-km)", "Millones ton-km", "Parte 2: Infraestructura (Transporte)"),
        "IE.PPI.TRAN.CD": ("Inversión privada en transporte (US$ corrientes)", "Investment in transport with private participation (current US$)", "US$", "Parte 2: Infraestructura (Transporte)"),
        "EG.ELC.PROD.KH": ("Producción de electricidad total (kWh)", "Electricity production (kWh)", "kWh", "Parte 2: Infraestructura (Energía)"),
        "EG.ELC.COAL.ZS": ("Producción de electricidad a partir de carbón (%)", "Electricity production from coal sources (% of total)", "%", "Parte 2: Infraestructura (Energía)"),
        "EG.ELC.HYRO.ZS": ("Producción de electricidad de origen hidroeléctrico (%)", "Electricity production from hydroelectric sources (% of total)", "%", "Parte 2: Infraestructura (Energía)"),
        "EG.USE.COMM.FO.ZS": ("Consumo de energía fósil (% del total)", "Fossil fuel energy consumption (% of total)", "%", "Parte 2: Infraestructura (Energía)"),
        "IE.PPI.ENGY.CD": ("Inversión privada en energía (US$ corrientes)", "Investment in energy with private participation (current US$)", "US$", "Parte 2: Infraestructura (Energía)"),
        "IT.MLT.MAIN.P2": ("Líneas de telefonía fija (por 100 personas)", "Telephone lines (per 100 people)", "Por 100 hab.", "Parte 2: Infraestructura (Comunicaciones)"),
        "IT.TEL.INVS.RV.ZS": ("Inversión en telecomunicaciones (% de los ingresos)", "Telecommunications investment (% of revenue)", "%", "Parte 2: Infraestructura (Comunicaciones)"),
        "IE.PPI.TELE.CD": ("Inversión privada en telecomunicaciones (US$ corrientes)", "Investment in telecoms with private participation (current US$)", "US$", "Parte 2: Infraestructura (Comunicaciones)")
    }
    
    for code, col_name in econstats_infra_vars.items():
        desc, desc_en, unit, part = econ_meta_desc[code]
        valid = df_part2.dropna(subset=[col_name])
        cob_str = f"{valid['year'].min()} - {valid['year'].max()} (n={len(valid)})" if not valid.empty else "Sin datos"
        metadata_rows.append({
            "Bloque": part,
            "Columna": col_name,
            "Nombre Variable (Español)": desc,
            "Nombre Original (Inglés)": desc_en,
            "Código Indicador": code,
            "Fuente Oficial": "EconStats (Historical Series IND)",
            "Unidad de Medida": unit,
            "Cobertura": cob_str
        })
        
    df_meta = pd.DataFrame(metadata_rows)
    
    # Escribir con openpyxl aplicando estilos
    for target_excel in [excel_path_proc, excel_path_out]:
        with pd.ExcelWriter(target_excel, engine="openpyxl") as writer:
            df_part1.to_excel(writer, sheet_name="Crecimiento_e_Inversion", index=False)
            df_part2.to_excel(writer, sheet_name="Infraestructura", index=False)
            df_panel.to_excel(writer, sheet_name="Panel_Consolidado", index=False)
            df_meta.to_excel(writer, sheet_name="Diccionario_Variables", index=False)
            
            # Formato y ajuste de anchos
            wb = writer.book
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            header_fill_blue = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_fill_green = PatternFill(start_color="1E4620", end_color="1E4620", fill_type="solid")
            header_fill_gray = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
            
            fills = {
                "Crecimiento_e_Inversion": header_fill_blue,
                "Infraestructura": header_fill_green,
                "Panel_Consolidado": header_fill_gray,
                "Diccionario_Variables": header_fill_blue
            }
            
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                ws.views.sheetView[0].showGridLines = True
                
                # Formatear encabezado
                fill_to_use = fills.get(sheet_name, header_fill_blue)
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = fill_to_use
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border
                ws.row_dimensions[1].height = 28
                
                # Ajustar ancho de columnas y formato numérico
                for col in ws.columns:
                    col_letter = col[0].column_letter
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
                    
                    # Formatear filas de datos
                    for cell in col[1:]:
                        cell.border = thin_border
                        if isinstance(cell.value, float):
                            cell.number_format = "#,##0.00"
                        elif isinstance(cell.value, int) and col_letter == 'C': # Columna year
                            cell.number_format = "0"
                            
        print(f"   [Guardado Excel]: {target_excel}")

    print("\n" + "=" * 80)
    print(" ¡PROCESO DE DESCARGA Y EXPORTACIÓN FINALIZADO CON ÉXITO!")
    print("=" * 80)

if __name__ == "__main__":
    main()
