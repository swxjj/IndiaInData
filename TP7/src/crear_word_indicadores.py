"""Crea una selección breve de indicadores para el punto 2."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'tp7_punto2_indicadores.docx'
OUT.parent.mkdir(exist_ok=True, parents=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08
for name in ['Title', 'Subtitle', 'Heading 1']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(6)
doc.add_paragraph('Indicadores laborales para India', 'Title')
doc.add_paragraph('TP7 · Punto 2 · Trabajo y desarrollo', 'Subtitle')
doc.add_paragraph('Estos seis indicadores permiten mirar cuántas personas participan del mercado laboral, qué oportunidades tienen y qué calidad de empleo consiguen.')

items = [
    ('Tasa de participación laboral',
     'Mide qué proporción de la población en edad de trabajar tiene empleo o lo busca. Comparar mujeres y varones permite detectar diferencias en el acceso al mercado laboral.'),
    ('Relación entre ocupación y población',
     'Muestra qué proporción de la población en edad de trabajar está ocupada. Sirve para ver cuánto empleo genera la economía, aunque no informa por sí sola sobre su calidad.'),
    ('Tasa de desempleo',
     'Indica qué parte de la fuerza de trabajo busca empleo y no lo consigue. Es útil para observar dificultades de inserción, pero debe leerse junto con la participación y la informalidad.'),
    ('Tasa de ocupación informal',
     'Permite observar cuánto empleo carece de registro o protección formal. Es interesante porque tener trabajo no siempre implica contar con protección social y derechos laborales efectivos.'),
    ('Jóvenes que no estudian ni trabajan ni reciben formación',
     'La tasa NEET se refiere a jóvenes de 15 a 24 años. Ayuda a identificar dificultades en el paso de la educación al empleo y riesgos de quedar fuera de ambos ámbitos.'),
    ('Tasa de pobreza laboral',
     'Mide la proporción de personas ocupadas que viven en hogares en pobreza. Permite preguntarse si el empleo alcanza para sostener un nivel de vida básico.'),
]
for title, body in items:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    p.add_run(title + '. ').bold = True
    p.add_run(body)

p = doc.add_paragraph('Al relevar los datos, conviene anotar el año de cada indicador y los valores de mujeres y varones cuando estén disponibles. El perfil puede combinar años diferentes.')
p.paragraph_format.space_before = Pt(4)

def source(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run(label + ': ').font.size = Pt(9)
    link = OxmlElement('w:hyperlink')
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); props.append(sz)
    run.append(props)
    text = OxmlElement('w:t'); text.text = url; run.append(text)
    link.append(run); p._p.append(link)

source('Fuente principal, OIT ILOSTAT, perfil de India', 'https://ilostat.ilo.org/es/data/country-profiles/ind/')
source('Referencia sobre informalidad, OIT', 'https://www.ilo.org/resource/45-informal-economy-workers')
p = doc.add_paragraph('Consulta: 21 de octubre de 2026.')
p.runs[0].font.size = Pt(9)
doc.core_properties.title = 'Indicadores laborales para India'
doc.core_properties.subject = 'TP7 Punto 2 Selección de indicadores'
doc.core_properties.author = ''
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
